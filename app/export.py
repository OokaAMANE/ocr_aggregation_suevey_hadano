"""集計結果・回答データの Excel / CSV 出力。"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import aggregate
from .surveys import CHOICE_TYPES

STATUS_LABELS = {
    "queued": "読取待ち",
    "processing": "読取中",
    "auto": "読取済（未確認）",
    "review": "要確認",
    "confirmed": "確認済",
    "error": "エラー",
}

HEAD_FILL = PatternFill("solid", fgColor="DDE7F3")
TITLE_FONT = Font(bold=True, size=12)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------- 回答データ（1行 = 1人）


def raw_columns(survey: dict) -> list[tuple[str, callable]]:
    cols: list[tuple[str, callable]] = []
    for q in survey["questions"]:
        qid, no, t = q["id"], q["no"], q["type"]
        if t == "single":
            cols.append((no, lambda a, qid=qid: "／".join(a[qid]["selected"])))
            if q.get("other_option"):
                cols.append((f"{no}_その他記述", lambda a, qid=qid: a[qid]["other_text"]))
        elif t == "multi":
            for o in q["options"]:
                cols.append((f"{no}_{o}", lambda a, qid=qid, o=o: 1 if o in a[qid]["selected"] else ""))
            if q.get("other_option"):
                cols.append((f"{no}_その他記述", lambda a, qid=qid: a[qid]["other_text"]))
        elif t == "timegrid":
            for row in q["rows"]:
                for h in q["hours"]:
                    cols.append(
                        (f"{no}_{row['label']}_{h}時",
                         lambda a, qid=qid, rid=row["id"], h=h: 1 if h in a[qid][rid] else "")
                    )
        elif t == "destinations":
            for p in q["purposes"]:
                pid, pl = p["id"], p["label"]
                cols.append((f"{no}_{pl}", lambda a, qid=qid, pid=pid: 1 if a[qid][pid]["checked"] else ""))
                cols.append((f"{no}_{pl}_目的地", lambda a, qid=qid, pid=pid: a[qid][pid]["place"]))
                cols.append((f"{no}_{pl}_停留所", lambda a, qid=qid, pid=pid: a[qid][pid]["stop"]))
        elif t == "trip":
            cols.append((f"{no}_目的地", lambda a, qid=qid: a[qid]["destination"]))
            cols.append((f"{no}_出発時刻", lambda a, qid=qid: a[qid]["depart_time"]))
            cols.append((f"{no}_帰りの到着時刻", lambda a, qid=qid: a[qid]["return_time"]))
        else:
            cols.append((no, lambda a, qid=qid: a[qid]["text"]))
    return cols


def raw_rows(survey: dict, responses: list[dict]) -> tuple[list[str], list[list]]:
    cols = raw_columns(survey)
    header = ["No.", "状態", "回収区分", "ファイル", *[c[0] for c in cols]]
    rows = [
        [r["id"], STATUS_LABELS.get(r["status"], r["status"]), r["batch"], r["source"],
         *[f(r["answers"]) for _, f in cols]]
        for r in responses
    ]
    return header, rows


def to_csv(survey: dict, responses: list[dict]) -> bytes:
    header, rows = raw_rows(survey, responses)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # Excel で文字化けしないよう BOM 付き


# ---------------------------------------------------------------- Excel


def _cell(ws, row, col, value, **style):
    c = ws.cell(row=row, column=col, value=value)
    for k, v in style.items():
        setattr(c, k, v)
    return c


def _write_simple(ws, summary: dict, title: str):
    _cell(ws, 1, 1, title, font=TITLE_FONT)
    _cell(ws, 2, 1, f"集計対象: {summary['total']}人")
    r = 4
    for q in summary["questions"]:
        _cell(ws, r, 1, f"{q['no']}　{q['label']}", font=BOLD)
        r += 1
        t = q["type"]
        if t in CHOICE_TYPES:
            for i, h in enumerate(["選択肢", "回答数", "割合(%)"], 1):
                _cell(ws, r, i, h, fill=HEAD_FILL, border=BOX, font=BOLD)
            r += 1
            for row in q["rows"]:
                if row.get("extra") and not row["count"]:
                    continue
                _cell(ws, r, 1, row["label"], border=BOX)
                _cell(ws, r, 2, row["count"], border=BOX)
                _cell(ws, r, 3, row["pct"], border=BOX)
                r += 1
            _cell(ws, r, 1, "回答者数（母数）", border=BOX, font=BOLD)
            _cell(ws, r, 2, q["total"], border=BOX, font=BOLD)
            r += 1
            if q["type"] == "multi":
                _cell(ws, r, 1, "※複数回答のため割合の合計は100%を超えることがあります")
                r += 1
            for txt in q["other_texts"]:
                _cell(ws, r, 1, f"　その他記述: {txt}")
                r += 1
        elif t == "timegrid":
            _cell(ws, r, 1, "", fill=HEAD_FILL, border=BOX)
            for i, h in enumerate(q["hours"], 2):
                _cell(ws, r, i, f"{h}時", fill=HEAD_FILL, border=BOX, font=BOLD)
            r += 1
            for row in q["rows"]:
                _cell(ws, r, 1, row["label"], border=BOX)
                for i, n in enumerate(row["counts"], 2):
                    _cell(ws, r, i, n, border=BOX)
                r += 1
        elif t == "destinations":
            for i, h in enumerate(["種別", "回答数", "割合(%)", "主な目的地（件数）", "主な最寄り停留所（件数）"], 1):
                _cell(ws, r, i, h, fill=HEAD_FILL, border=BOX, font=BOLD)
            r += 1
            for p in q["purposes"]:
                _cell(ws, r, 1, p["label"], border=BOX)
                _cell(ws, r, 2, p["count"], border=BOX)
                _cell(ws, r, 3, p["pct"], border=BOX)
                _cell(ws, r, 4, "、".join(f"{x['label']}({x['count']})" for x in p["places"]), border=BOX)
                _cell(ws, r, 5, "、".join(f"{x['label']}({x['count']})" for x in p["stops"]), border=BOX)
                r += 1
        elif t == "trip":
            for label, key in (("目的地", "destinations"), ("出発時刻（時台）", "depart_hours"),
                               ("帰りの到着時刻（時台）", "return_hours")):
                _cell(ws, r, 1, label, fill=HEAD_FILL, border=BOX, font=BOLD)
                _cell(ws, r, 2, "件数", fill=HEAD_FILL, border=BOX, font=BOLD)
                r += 1
                for x in q[key]:
                    _cell(ws, r, 1, x["label"], border=BOX)
                    _cell(ws, r, 2, x["count"], border=BOX)
                    r += 1
        elif t == "short_text":
            for i, h in enumerate(["記入内容", "件数"], 1):
                _cell(ws, r, i, h, fill=HEAD_FILL, border=BOX, font=BOLD)
            r += 1
            for x in q["values"]:
                _cell(ws, r, 1, x["label"], border=BOX)
                _cell(ws, r, 2, x["count"], border=BOX)
                r += 1
        else:
            _cell(ws, r, 1, f"記入あり {len(q['texts'])}件（「自由記述」シート参照）")
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 44
    for col in "BCDE":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 40


def _write_cross(ws, survey: dict, responses: list[dict], by_q: dict):
    _cell(ws, 1, 1, f"クロス集計（{by_q['label']}別）", font=TITLE_FONT)
    _cell(ws, 2, 1, "上段: 人数　下段: 割合(%)　割合の母数は各行の回答者数")
    r = 4
    for q in survey["questions"]:
        if q["type"] not in CHOICE_TYPES or q["id"] == by_q["id"]:
            continue
        ct = aggregate.cross_table(survey, by_q["id"], q, responses)
        _cell(ws, r, 1, f"{q['no']}　{q['label']}", font=BOLD)
        r += 1
        _cell(ws, r, 1, by_q["label"], fill=HEAD_FILL, border=BOX, font=BOLD)
        _cell(ws, r, 2, "回答者数", fill=HEAD_FILL, border=BOX, font=BOLD)
        for i, c in enumerate(ct["columns"], 3):
            _cell(ws, r, i, c, fill=HEAD_FILL, border=BOX, font=BOLD,
                  alignment=Alignment(wrap_text=True, vertical="top"))
        r += 1
        for row in ct["rows"]:
            _cell(ws, r, 1, row["label"], border=BOX)
            _cell(ws, r, 2, row["n"], border=BOX)
            for i, n in enumerate(row["counts"], 3):
                _cell(ws, r, i, n, border=BOX)
            r += 1
            _cell(ws, r, 1, "", border=BOX)
            _cell(ws, r, 2, "", border=BOX)
            for i, p in enumerate(row["pcts"], 3):
                _cell(ws, r, i, p, border=BOX, font=Font(color="666666"))
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 18
    for i in range(2, 20):
        ws.column_dimensions[get_column_letter(i)].width = 13


def to_xlsx(survey: dict, responses: list[dict], filter_label: str = "") -> bytes:
    wb = Workbook()
    summary = aggregate.summarize(survey, responses)
    ws = wb.active
    ws.title = "単純集計"
    _write_simple(ws, summary, survey["title"] + (f"（{filter_label}）" if filter_label else ""))

    for q in survey["questions"]:
        if q.get("attribute"):
            _write_cross(wb.create_sheet(f"クロス_{q['no']}"), survey, responses, q)

    ws = wb.create_sheet("自由記述")
    _cell(ws, 1, 1, "No.", fill=HEAD_FILL, font=BOLD)
    _cell(ws, 1, 2, "設問", fill=HEAD_FILL, font=BOLD)
    _cell(ws, 1, 3, "記述内容", fill=HEAD_FILL, font=BOLD)
    r = 2
    for q in summary["questions"]:
        for t in q.get("texts", []):
            _cell(ws, r, 1, t["id"])
            _cell(ws, r, 2, q["no"])
            _cell(ws, r, 3, t["text"], alignment=Alignment(wrap_text=True, vertical="top"))
            r += 1
    ws.column_dimensions["C"].width = 100

    ws = wb.create_sheet("回答データ")
    header, rows = raw_rows(survey, responses)
    ws.append(header)
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = BOLD
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
