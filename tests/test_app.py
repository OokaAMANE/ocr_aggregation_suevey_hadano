import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

import json
from types import SimpleNamespace as NS

from app import aggregate, jobs, main, ocr, surveys


class InlineExecutor:
    """読み取りをバックグラウンドではなくその場で実行する（テスト用）。"""

    def submit(self, fn, *args):
        fn(*args)

    def shutdown(self, **kw):
        pass


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SURVEY_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(jobs, "_executor", InlineExecutor())
    with TestClient(main.app) as c:
        yield c


def png(color="white") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (400, 560), color).save(buf, "PNG")
    return buf.getvalue()


def pdf(pages: int) -> bytes:
    ims = [Image.new("RGB", (400, 560), "white") for _ in range(pages)]
    buf = io.BytesIO()
    ims[0].save(buf, "PDF", save_all=True, append_images=ims[1:])
    return buf.getvalue()


def answers_for(survey_id: str, **overrides) -> dict:
    a = surveys.blank_answers(surveys.get_survey(survey_id))
    a.update(overrides)
    return a


# ---------------------------------------------------------------- 様式定義


@pytest.mark.parametrize("survey", surveys.list_surveys(), ids=lambda s: s["id"])
def test_survey_definitions_are_consistent(survey):
    ids = [q["id"] for q in survey["questions"]]
    assert len(ids) == len(set(ids))
    for q in survey["questions"]:
        if q["type"] in surveys.CHOICE_TYPES:
            assert len(q["options"]) == len(set(q["options"]))
            if q.get("other_option"):
                assert q["other_option"] in q["options"]
    schema = surveys.ocr_schema(survey)
    assert schema["properties"]["answers"]["required"] == ids
    # 空の回答は正規化しても変わらない
    blank = surveys.blank_answers(survey)
    assert surveys.normalize_answers(survey, blank) == blank


def test_normalize_drops_unknown_values_and_fixes_formats():
    s = surveys.get_survey("taxi_minami")
    raw = {
        "gender": {"selected": ["女性", "宇宙人"], "other_text": ""},
        "age": {"selected": "７０代"},  # 全角・文字列でも受け付ける
        "q3": {"outbound": [8, "9", 99], "return": []},
        "q4": {"hospital": {"checked": False, "place": "秦野赤十字病院", "stop": ""}},
        "q6": {"destination": "Ａスーパー", "depart_time": "9時5", "return_time": "25:00"},
    }
    a = surveys.normalize_answers(s, raw)
    assert a["gender"]["selected"] == ["女性"]
    assert a["age"]["selected"] == ["70代"]
    assert a["q3"] == {"outbound": [8, 9], "return": []}
    assert a["q4"]["hospital"] == {"checked": True, "place": "秦野赤十字病院", "stop": ""}
    assert a["q6"] == {"destination": "Aスーパー", "depart_time": "09:05", "return_time": ""}
    assert surveys.norm_time("午後1時") == "13:00"
    assert surveys.norm_time("午前11:30") == "11:30"


# ---------------------------------------------------------------- 取り込み


def test_upload_splits_pdf_pages_into_respondents(client):
    files = [("files", ("scan.pdf", pdf(4), "application/pdf"))]
    res = client.post(
        "/api/surveys/chimura_new/upload",
        files=files,
        data={"batch": "第1回", "pages_per_response": "2", "ocr_mode": "none"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["pages"] == 4 and len(body["created"]) == 2 and body["warning"] == ""

    r = client.get(f"/api/responses/{body['created'][1]}").json()
    assert [p["page_no"] for p in r["pages"]] == [1, 2]
    assert r["source"] == "scan.pdf p.3,4"
    assert r["batch"] == "第1回"
    img = client.get(f"/api/pages/{r['pages'][0]['id']}/image")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"


def test_upload_groups_pages_within_each_pdf(client):
    # 1つ目の PDF は1ページ抜けている。2つ目の PDF の区切りはずれない
    files = [
        ("files", ("box01.pdf", pdf(5), "application/pdf")),
        ("files", ("box02.pdf", pdf(4), "application/pdf")),
    ]
    body = client.post(
        "/api/surveys/chimura_new/upload",
        files=files,
        data={"pages_per_response": "2", "ocr_mode": "none"},
    ).json()
    assert body["pages"] == 9 and len(body["created"]) == 5
    assert "box01.pdf" in body["warning"] and "box02.pdf" not in body["warning"]
    sources = [client.get(f"/api/responses/{i}").json()["source"] for i in body["created"]]
    assert sources == ["box01.pdf p.1,2", "box01.pdf p.3,4", "box01.pdf p.5", "box02.pdf p.1,2", "box02.pdf p.3,4"]


def test_upload_pairs_single_page_images_in_order(client):
    files = [("files", (n, png(), "image/png")) for n in ("1表.png", "1裏.png", "2表.png", "2裏.png")]
    body = client.post(
        "/api/surveys/chimura_new/upload",
        files=files,
        data={"pages_per_response": "2", "ocr_mode": "none"},
    ).json()
    sources = [client.get(f"/api/responses/{i}").json()["source"] for i in body["created"]]
    assert sources == ["1表.png / 1裏.png", "2表.png / 2裏.png"] and body["warning"] == ""


def test_upload_rejects_ocr_without_api_key(client, monkeypatch):
    monkeypatch.setattr(ocr, "available", lambda: False)
    res = client.post(
        "/api/surveys/taxi_oone/upload",
        files=[("files", ("a.png", png(), "image/png"))],
        data={"ocr_mode": "now"},
    )
    assert res.status_code == 400


def test_ocr_result_sets_status(client, monkeypatch):
    monkeypatch.setattr(ocr, "available", lambda: True)

    def fake_read(survey, paths):
        assert len(paths) == 1 and paths[0].exists()
        a = answers_for(survey["id"])
        a["gender"] = {"selected": ["男性"], "other_text": ""}
        return {"answers": a, "uncertain": [], "notes": ""}

    monkeypatch.setattr(ocr, "read_form", fake_read)
    res = client.post(
        "/api/surveys/taxi_oone/upload",
        files=[("files", ("a.png", png(), "image/png"))],
        data={"ocr_mode": "now"},
    ).json()
    r = client.get(f"/api/responses/{res['created'][0]}").json()
    assert r["status"] == "auto"
    assert r["answers"]["gender"]["selected"] == ["男性"]

    # 単一回答に複数チェック → 要確認
    def fake_read_multi(survey, paths):
        a = answers_for(survey["id"])
        a["age"] = {"selected": ["60代", "70代"], "other_text": ""}
        return {"answers": a, "uncertain": [], "notes": ""}

    monkeypatch.setattr(ocr, "read_form", fake_read_multi)
    client.post(f"/api/responses/{res['created'][0]}/ocr")
    r = client.get(f"/api/responses/{res['created'][0]}").json()
    assert r["status"] == "review" and r["uncertain"] == ["age"]

    # 読み取り失敗はエラーとして記録される
    def fake_fail(survey, paths):
        raise ocr.OCRError("だめでした")

    monkeypatch.setattr(ocr, "read_form", fake_fail)
    client.post(f"/api/responses/{res['created'][0]}/ocr")
    r = client.get(f"/api/responses/{res['created'][0]}").json()
    assert r["status"] == "error" and r["error"] == "だめでした"


# ---------------------------------------------------------------- 確認・集計・出力


def add_manual(client, survey_id, answers, status="confirmed"):
    rid = client.post(f"/api/surveys/{survey_id}/responses").json()["id"]
    res = client.put(f"/api/responses/{rid}", json={"answers": answers, "status": status})
    assert res.status_code == 200, res.text
    return rid


def sel(*values, other=""):
    return {"selected": list(values), "other_text": other}


def test_summary_cross_filter_and_export(client):
    sid = "taxi_nishi_chimura"
    add_manual(client, sid, answers_for(
        sid, gender=sel("男性"), age=sel("70代"), area=sel("千村台"), q2=sel("ほぼ毎日"),
        q3={"outbound": [8, 9], "return": [12]},
        q4={"hospital": {"checked": True, "place": "秦野赤十字病院", "stop": "日赤前"}},
        q5=sel("行きたい目的地がない", "その他", other="坂が多い"),
        q8={"text": "本数を増やしてほしい"},
    ))
    add_manual(client, sid, answers_for(
        sid, gender=sel("女性"), age=sel("70代"), q2=sel("ほぼ毎日", "月に数回程度"),
        q3={"outbound": [8], "return": []},
        q4={"hospital": {"checked": True, "place": "秦野赤十字病院", "stop": ""},
            "shopping": {"checked": True, "place": "", "stop": ""}},
        q6={"destination": "渋沢駅", "depart_time": "13:30", "return_time": ""},
    ))
    add_manual(client, sid, answers_for(sid, gender=sel("女性"), q2=sel("年に数回程度")), status="review")

    s = client.get(f"/api/surveys/{sid}/summary").json()
    assert s["total"] == 3
    q = {x["id"]: x for x in s["questions"]}
    counts = {r["label"]: r["count"] for r in q["q2"]["rows"]}
    assert counts["ほぼ毎日"] == 1 and counts["年に数回程度"] == 1
    assert counts[aggregate.INVALID] == 1  # 単一回答への複数チェック
    assert q["q3"]["rows"][0]["counts"][q["q3"]["hours"].index(8)] == 2
    hosp = q["q4"]["purposes"][0]
    assert hosp["count"] == 2 and hosp["places"] == [{"label": "秦野赤十字病院", "count": 2}]
    assert q["q5"]["other_texts"] == ["坂が多い"]
    assert q["q6"]["depart_hours"] == [{"label": "13時台", "hour": 13, "count": 1}]
    assert q["q8"]["texts"][0]["text"] == "本数を増やしてほしい"

    # 確認済のみ・絞り込み・クロス集計
    assert client.get(f"/api/surveys/{sid}/summary?scope=confirmed").json()["total"] == 2
    assert client.get(f"/api/surveys/{sid}/summary?f_gender=女性").json()["total"] == 2
    s = client.get(f"/api/surveys/{sid}/summary?cross=gender").json()
    ct = next(x for x in s["questions"] if x["id"] == "q2")["cross"]
    male = next(r for r in ct["rows"] if r["label"] == "男性")
    assert male["n"] == 1 and male["counts"][ct["columns"].index("ほぼ毎日")] == 1

    xlsx = client.get(f"/api/surveys/{sid}/export.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["単純集計", "クロス_性別", "クロス_年齢", "クロス_自治会区", "自由記述", "回答データ"]
    raw = list(wb["回答データ"].values)
    assert len(raw) == 4 and "問5_その他記述" in raw[0]

    csv_text = client.get(f"/api/surveys/{sid}/export.csv").content.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("No.,状態,回収区分,ファイル,性別")


def test_confirm_auto_and_delete(client, monkeypatch):
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(ocr, "read_form", lambda s, p: {"answers": answers_for(s["id"]), "uncertain": [], "notes": ""})
    created = client.post(
        "/api/surveys/taxi_minami/upload",
        files=[("files", ("a.png", png(), "image/png")), ("files", ("b.png", png(), "image/png"))],
        data={"ocr_mode": "now"},
    ).json()["created"]
    assert client.post("/api/surveys/taxi_minami/confirm-auto").json() == {"confirmed": 2}
    assert client.delete(f"/api/responses/{created[0]}").json() == {"ok": True}
    counts = client.get("/api/surveys/taxi_minami/responses").json()["counts"]
    assert counts["confirmed"] == 1


# ---------------------------------------------------------------- バッチ読み取り


class FakeBatchAPI:
    """client.messages.batches の代わり。送信内容を記録し、指定した結果を返す。"""

    def __init__(self):
        self.created: dict[str, list[dict]] = {}
        self.ended: set[str] = set()
        self.results_by_id: dict[str, dict] = {}  # custom_id -> 結果

    def create(self, requests):
        bid = f"msgbatch_{len(self.created) + 1}"
        self.created[bid] = list(requests)
        return NS(id=bid)

    def retrieve(self, bid):
        n = len(self.created[bid])
        done = bid in self.ended
        return NS(
            processing_status="ended" if done else "in_progress",
            request_counts=NS(processing=0 if done else n, succeeded=n if done else 0, errored=0, canceled=0, expired=0),
        )

    def results(self, bid):
        for req in self.created[bid]:
            cid = req["custom_id"]
            if cid in self.results_by_id:
                yield NS(custom_id=cid, result=self.results_by_id[cid])


def ok_result(survey_id, **answers):
    data = {"answers": answers_for(survey_id, **answers), "uncertain": [], "notes": ""}
    msg = NS(stop_reason="end_turn", content=[NS(type="text", text=json.dumps(data, ensure_ascii=False))])
    return NS(type="succeeded", message=msg)


@pytest.fixture
def fake_batches(monkeypatch):
    api = FakeBatchAPI()
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(ocr, "_get_client", lambda: NS(messages=NS(batches=api)))
    return api


def test_batch_ocr_submits_polls_and_stores_results(client, fake_batches, monkeypatch):
    sid = "taxi_minami"
    created = client.post(
        f"/api/surveys/{sid}/upload",
        files=[("files", (f"{i}.png", png(), "image/png")) for i in range(3)],
        data={"ocr_mode": "batch"},
    ).json()["created"]
    assert client.get(f"/api/surveys/{sid}/responses").json()["counts"]["batch"] == 3

    jobs.submit_pending()
    (bid, reqs), = fake_batches.created.items()
    assert [r["custom_id"] for r in reqs] == [f"r{i}" for i in created]
    params = reqs[0]["params"]
    assert params["output_config"]["format"]["type"] == "json_schema"
    assert "fallbacks" not in params  # バッチでは使えない
    assert [c["type"] for c in params["messages"][0]["content"]] == ["text", "image", "text"]
    jobs.submit_pending()  # 送信済みのものは二重に送らない
    assert len(fake_batches.created) == 1

    jobs.poll_batches()  # まだ処理中
    assert client.get(f"/api/surveys/{sid}/batches").json()["batches"][0]["processing"] == 3

    # 1件目は成功、2件目は失敗、3件目は結果待ちの間に通常読み取りでやり直された
    fake_batches.results_by_id = {
        f"r{created[0]}": ok_result(sid, gender=sel("女性")),
        f"r{created[1]}": NS(type="errored", error=NS(error=NS(type="overloaded_error", message="Overloaded"))),
        f"r{created[2]}": ok_result(sid, gender=sel("男性")),
    }
    monkeypatch.setattr(ocr, "read_form", lambda s, p: {"answers": answers_for(s["id"]), "uncertain": ["q1"], "notes": ""})
    client.post(f"/api/responses/{created[2]}/ocr")
    fake_batches.ended.add(bid)
    jobs.poll_batches()

    r0, r1, r2 = (client.get(f"/api/responses/{i}").json() for i in created)
    assert r0["status"] == "auto" and r0["answers"]["gender"]["selected"] == ["女性"]
    assert r1["status"] == "error" and "Overloaded" in r1["error"]
    assert r2["status"] == "review" and r2["answers"]["gender"]["selected"] == []  # 古いバッチ結果で上書きしない
    assert client.get(f"/api/surveys/{sid}/batches").json()["batches"][0]["status"] == "ended"

    # エラーはバッチで再読み取りできる
    assert client.post(f"/api/surveys/{sid}/retry-errors?mode=batch").json() == {"queued": 1}
    jobs.submit_pending()
    assert len(fake_batches.created) == 2


def test_batch_is_split_by_size(client, fake_batches, monkeypatch):
    monkeypatch.setattr(jobs, "MAX_BATCH_BYTES", 1)  # 1件ずつ別のバッチになる大きさ
    client.post(
        "/api/surveys/taxi_minami/upload",
        files=[("files", (f"{i}.png", png(), "image/png")) for i in range(3)],
        data={"ocr_mode": "batch"},
    )
    jobs.submit_pending()
    assert [len(r) for r in fake_batches.created.values()] == [1, 1, 1]
